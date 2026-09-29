"""Rough text-fit check for a .pptx when no renderer is available.

Estimates wrapped line count from average glyph width and flags text boxes whose
text needs more height than the box has, or shapes that leave the slide.

    python tools/fitcheck.py deck.pptx
"""

import math
import sys

from pptx import Presentation
from pptx.util import Emu

AVG_W = {"Calibri": 0.50, "Cambria": 0.54}  # average glyph width as a fraction of font size


def check(path):
    prs = Presentation(path)
    sw, sh = prs.slide_width, prs.slide_height
    problems = 0
    for i, slide in enumerate(prs.slides, 1):
        for shp in slide.shapes:
            if shp.left is not None and (shp.left + shp.width > sw + 10 or shp.top + shp.height > sh + 10):
                print(f"slide {i}: '{shp.name}' extends past slide edge")
                problems += 1
            if not shp.has_text_frame or not shp.text_frame.text.strip():
                continue
            w_pt = Emu(shp.width).pt
            h_pt = Emu(shp.height).pt
            need = 0.0
            for p in shp.text_frame.paragraphs:
                text = "".join(r.text for r in p.runs)
                size = max((r.font.size.pt for r in p.runs if r.font.size), default=18)
                face = next((r.font.name for r in p.runs if r.font.name), "Calibri")
                bullet_indent = 20 if p._pPr is not None and p._pPr.find(
                    "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None else 0
                cw = size * AVG_W.get(face, 0.52) * (1.08 if any(r.font.bold for r in p.runs) else 1)
                per_line = max(1, int((w_pt - bullet_indent) / cw))
                lines = max(1, math.ceil(len(text) / per_line)) if text else 1
                need += lines * size * 1.2 + 6
            if need > h_pt * 1.05:
                print(f"slide {i}: text may overflow ({need:.0f}pt needed, {h_pt:.0f}pt box): "
                      f"{shp.text_frame.text[:60]!r}")
                problems += 1
    print(f"{problems} possible problem(s)")


if __name__ == "__main__":
    check(sys.argv[1])
