"""Approximate HTML preview of a .pptx (shapes, images, text) for visual QA when
LibreOffice is unavailable. Charts are drawn as labelled placeholders.

    python tools/preview_pptx.py deck.pptx out.html
"""

import base64
import html
import sys

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu

PX = 96  # pixels per inch


def px(v):
    return Emu(v).inches * PX


def color_of(fill):
    try:
        if fill.type == 1:
            return "#" + str(fill.fore_color.rgb)
    except Exception:
        pass
    return None


def render(path, out):
    prs = Presentation(path)
    W, H = px(prs.slide_width), px(prs.slide_height)
    parts = ["<html><body style='background:#666;margin:0;font-family:Calibri,Carlito,sans-serif'>"]
    for idx, slide in enumerate(prs.slides, 1):
        bg = "#FFFFFF"
        try:
            bg = "#" + str(slide.background.fill.fore_color.rgb)
        except Exception:
            pass
        parts.append(f"<div class=slide id=s{idx} style='position:relative;width:{W}px;height:{H}px;"
                     f"background:{bg};margin:20px;overflow:hidden'>")
        for shp in slide.shapes:
            x, y, w, h = px(shp.left), px(shp.top), px(shp.width), px(shp.height)
            box = f"position:absolute;left:{x}px;top:{y}px;width:{w}px;height:{h}px;"
            if shp.shape_type == MSO_SHAPE_TYPE.PICTURE:
                b64 = base64.b64encode(shp.image.blob).decode()
                parts.append(f"<img style='{box}' src='data:{shp.image.content_type};base64,{b64}'>")
                continue
            if shp.has_chart:
                parts.append(f"<div style='{box}border:2px dashed #999;display:flex;align-items:center;"
                             f"justify-content:center;color:#666'>[native chart]</div>")
                continue
            style = box
            fillc = color_of(shp.fill) if hasattr(shp, "fill") else None
            if fillc:
                style += f"background:{fillc};"
            geom = getattr(shp, "auto_shape_type", None) if shp.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE else None
            if geom is not None and "OVAL" in str(geom):
                style += "border-radius:50%;"
            elif geom is not None and "ROUNDED" in str(geom):
                style += "border-radius:12px;"
            inner = ""
            if shp.has_text_frame and shp.text_frame.text.strip():
                va = {"MIDDLE (3)": "center", "BOTTOM (4)": "flex-end"}.get(str(shp.text_frame.vertical_anchor), "flex-start")
                ps = []
                for p in shp.text_frame.paragraphs:
                    align = {"CENTER (2)": "center", "RIGHT (3)": "right"}.get(str(p.alignment), "left")
                    bullet = "• " if p._pPr is not None and p._pPr.find(
                        "{http://schemas.openxmlformats.org/drawingml/2006/main}buChar") is not None else ""
                    spans = []
                    for r in p.runs:
                        f = r.font
                        c = "#" + str(f.color.rgb) if f.color and f.color.type is not None else "#000"
                        fam = "Cambria,Caladea,serif" if f.name == "Cambria" else "Calibri,Carlito,sans-serif"
                        spans.append(f"<span style='font-size:{f.size.pt if f.size else 18}pt;color:{c};font-family:{fam};"
                                     f"font-weight:{700 if f.bold else 400};font-style:{'italic' if f.italic else 'normal'}'>"
                                     f"{html.escape(r.text)}</span>")
                    pad = "padding-left:18px;text-indent:-14px;" if bullet else ""
                    ps.append(f"<div style='text-align:{align};margin-bottom:6px;line-height:1.15;{pad}'>{bullet}{''.join(spans)}</div>")
                inner = (f"<div style='display:flex;flex-direction:column;justify-content:{va};height:100%'>"
                         f"{''.join(ps)}</div>")
            parts.append(f"<div style=\"{style}\">{inner}</div>")
        parts.append("</div>")
    parts.append("</body></html>")
    open(out, "w").write("".join(parts))


if __name__ == "__main__":
    render(sys.argv[1], sys.argv[2])
